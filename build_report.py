"""
Step 6: Build the Excel management report.

- Sheet 1 "P&L Variance": clean variance table, flagged lines highlighted,
  Change and Change % as live Excel formulas so the finance team can check them.
- Sheet 2 "Commentary": the AI draft commentary from commentary.md.

Needs: pip install openpyxl
Run after pull_pnl.py and write_commentary.py.
"""

import json
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NAVY = "1F3864"
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")   # light yellow for review lines
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TOP_BORDER = Border(top=Side(style="thin", color="808080"))
MONEY = '#,##0.00;[Red]-#,##0.00'
PERCENT = '0.0%;[Red]-0.0%'


def build_variance_sheet(ws, data):
    ws.title = "P&L Variance"
    cur, prev = data["current_period"], data["previous_period"]

    ws["A1"] = f"{data['company']} - Profit and Loss Variance"
    ws["A1"].font = Font(size=14, bold=True, color=NAVY)
    ws["A2"] = f"{cur} vs {prev}"
    ws["A2"].font = Font(size=11, italic=True)
    ws["A3"] = ("Source: Xero API (read-only). Highlighted lines are above the review "
                "threshold. Draft for accountant review.")
    ws["A3"].font = Font(size=9, color="595959")

    headers = ["Section", "Account", cur, prev, "Change", "Change %", "Review"]
    header_row = 5
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center" if col > 2 else "left")

    r = header_row + 1
    for line in data["lines"]:
        ws.cell(row=r, column=1, value=line["section"])
        ws.cell(row=r, column=2, value=line["account"])
        ws.cell(row=r, column=3, value=line["current"]).number_format = MONEY
        ws.cell(row=r, column=4, value=line["previous"]).number_format = MONEY
        # Live formulas so the finance team can see how each number was made
        ws.cell(row=r, column=5, value=f"=C{r}-D{r}").number_format = MONEY
        ws.cell(row=r, column=6, value=(
            f'=IF(D{r}=0,IF(C{r}=0,"","new"),'
            f'IF(AND(D{r}<0,C{r}>0),"loss to profit",'
            f'IF(AND(D{r}>0,C{r}<0),"profit to loss",E{r}/ABS(D{r}))))'
        )).number_format = PERCENT
        ws.cell(row=r, column=6).alignment = Alignment(horizontal="right")
        ws.cell(row=r, column=7, value="Check" if line["flag"] else "")

        if line["flag"]:
            for col in range(1, 8):
                ws.cell(row=r, column=col).fill = FLAG_FILL
            ws.cell(row=r, column=7).font = Font(bold=True, color="C00000")
        if line["is_total"]:
            for col in range(1, 8):
                ws.cell(row=r, column=col).font = Font(bold=True)
                ws.cell(row=r, column=col).border = TOP_BORDER
        r += 1

    widths = [22, 32, 14, 14, 14, 14, 9]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = f"A{header_row + 1}"


def build_commentary_sheet(ws, commentary_md):
    ws.title = "Commentary"
    ws.column_dimensions["A"].width = 110
    r = 1
    for raw in commentary_md.splitlines():
        line = raw.strip()
        if not line or line == "---":
            continue
        text = re.sub(r"\*+", "", line)          # remove markdown bold/italic
        cell = ws.cell(row=r, column=1)
        if line.startswith("# "):
            cell.value = text[2:]
            cell.font = Font(size=14, bold=True, color=NAVY)
        elif line.startswith("## "):
            r += 1                                   # blank line before a section
            cell = ws.cell(row=r, column=1, value=text[3:])
            cell.font = Font(size=12, bold=True, color=NAVY)
        elif line.startswith("- "):
            cell.value = "\u2022 " + text[2:]
        else:
            cell.value = text
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1


def main():
    with open("variance.json") as f:
        data = json.load(f)
    with open("commentary.md", encoding="utf-8") as f:
        commentary_md = f.read()

    wb = Workbook()
    build_variance_sheet(wb.active, data)
    build_commentary_sheet(wb.create_sheet(), commentary_md)

    safe_period = re.sub(r"[^A-Za-z0-9]+", "_", data["current_period"]).strip("_")
    filename = f"Management_Report_{safe_period}.xlsx"
    wb.save(filename)
    print(f"Saved {filename}")


if __name__ == "__main__":
    main()
